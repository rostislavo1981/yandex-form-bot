"""Shared openpyxl styles. Defends against str(None).strip()=='None' bug."""
from __future__ import annotations

from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.styles.colors import Color

# Colors
HEADER_FILL = PatternFill(
    start_color=Color(rgb="FF1F4E78"),
    end_color=Color(rgb="FF1F4E78"),
    fill_type="solid",
)
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color=Color(rgb="FFFFFFFF"))

CONFIRMED_FILL = PatternFill(  # yellow: факт подтверждён скриншотом
    start_color=Color(rgb="FFFFF2CC"),
    end_color=Color(rgb="FFFFF2CC"),
    fill_type="solid",
)
ANOMALY_FILL = PatternFill(  # orange: >50% аномалия
    start_color=Color(rgb="FFFCE4D6"),
    end_color=Color(rgb="FFFCE4D6"),
    fill_type="solid",
)

THIN = Side(border_style="thin", color=Color(rgb="FFBFBFBF"))
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def safe_str(v: object) -> str:
    """Convert cell value to string, never producing 'None' or 'nan'."""
    if v is None:
        return ""
    s = str(v)
    if s in ("None", "nan", "NaN"):
        return ""
    return s


def safe_float(v: object) -> float:
    """Convert cell value to float, returning 0.0 for None/blank/non-numeric."""
    if v is None or v == "":
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0
