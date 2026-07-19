from __future__ import annotations

from io import BytesIO

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook
from sqlalchemy import select

from app.main import app
from app.models import Contract, Object, ObjectContract, ObjectStage, Stage
from app.services.excel_service import SHEET_COLUMNS, SHEET_ORDER

IMPORT_URL = "/api/catalogs/import"


@pytest.fixture(autouse=True)
def _admin_dev_user(db_session):
    """Import endpoints require manager/admin — pre-create dev-user as admin."""
    from app.models.users import User

    db_session.add(User(max_user_id="dev-user", full_name="Dev User", role="admin"))
    db_session.commit()


@pytest.fixture
async def client():
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
        yield c


def _build_workbook(rows_by_sheet: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for sheet_name in SHEET_ORDER:
        ws = wb.create_sheet(title=sheet_name)
        ws.append(SHEET_COLUMNS[sheet_name])
        for row in rows_by_sheet.get(sheet_name, []):
            ws.append([row.get(col) for col in SHEET_COLUMNS[sheet_name]])
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _build_simple_objects_workbook() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Список Ростислав"
    ws["B2"] = "Список в Яндекс Форму"
    ws["B4"] = "краткое название"
    ws["C4"] = "Договор"
    ws["B5"] = "Тамбасова"
    ws["C5"] = "РТ_26 Тамбасова — БКТП и КЛ"
    ws["D5"] = "СП_26 Тамбасова — реконструкция ТП"
    ws["B6"] = "Пискаревский"
    ws["C6"] = "РТ_26 Пискаревский — школа"
    ws["B7"] = "Обводный канал"
    ws["C7"] = "??"
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _build_object_only(obj_code: str = "obj-1") -> dict:
    return {
        "Contractors": [
            {"code": "c1", "name": "Подрядчик 1", "active": 1, "sort_order": 1},
        ],
        "Objects": [
            {
                "code": obj_code,
                "name": "Объект 1",
                "execution_method": "own",
                "default_contractor_code": "c1",
                "active": 1,
                "sort_order": 1,
            },
        ],
    }


def _object_mapping_row(
    obj_code: str,
    short_name: str,
    contract_code: str,
    full_name: str,
    primary: int = 1,
    active: int = 1,
) -> dict:
    return {
        "object_code": obj_code,
        "short_name": short_name,
        "contract_code": contract_code,
        "full_name": full_name,
        "primary": primary,
        "active": active,
    }


async def test_import_object_mappings_happy_path(client, db_session):
    rows = _build_object_only()
    rows["ObjectMappings"] = [
        _object_mapping_row("obj-1", "Объект 1 (раб)", "ctr-1", "Договор 1", 1, 1),
    ]
    content = _build_workbook(rows)

    validate_resp = await client.post(
        f"{IMPORT_URL}/validate",
        files={"file": ("catalogs.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert validate_resp.status_code == 200
    body = validate_resp.json()
    assert body["valid"] is True

    apply_resp = await client.post(
        f"{IMPORT_URL}/apply",
        files={"file": ("catalogs.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert apply_resp.status_code == 200
    assert apply_resp.json()["status"] == "applied"

    contract = db_session.execute(
        select(Contract).where(Contract.code == "ctr-1")
    ).scalar_one()
    assert contract.full_name == "Договор 1"

    link = db_session.execute(
        select(ObjectContract).where(ObjectContract.contract_id == contract.id)
    ).scalar_one()
    assert link.is_primary is True
    assert link.active is True


async def test_import_object_mappings_skip_empty_contract(client, db_session):
    rows = _build_object_only()
    rows["ObjectMappings"] = [
        _object_mapping_row("obj-1", "Объект 1 (раб)", "??", "", 1, 1),
        _object_mapping_row("obj-1", "Объект 1 (раб)", "нет пока договора", "", 1, 1),
    ]
    content = _build_workbook(rows)

    apply_resp = await client.post(
        f"{IMPORT_URL}/apply",
        files={"file": ("catalogs.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert apply_resp.status_code == 200

    assert db_session.execute(select(Contract)).scalar_one_or_none() is None
    assert db_session.execute(select(ObjectContract)).scalar_one_or_none() is None


async def test_import_object_mappings_idempotent(client, db_session):
    rows = _build_object_only()
    rows["ObjectMappings"] = [
        _object_mapping_row("obj-1", "Объект 1 (раб)", "ctr-1", "Договор 1", 1, 1),
        _object_mapping_row("obj-1", "Объект 1 (раб)", "ctr-2", "Договор 2", 0, 1),
    ]
    content = _build_workbook(rows)

    for _ in range(2):
        apply_resp = await client.post(
            f"{IMPORT_URL}/apply",
            files={"file": ("catalogs.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert apply_resp.status_code == 200

    assert len(db_session.execute(select(Contract)).scalars().all()) == 2
    assert len(db_session.execute(select(ObjectContract)).scalars().all()) == 2


async def test_import_object_mappings_rollback_on_error(client, db_session):
    rows = _build_object_only()
    rows["ObjectMappings"] = [
        _object_mapping_row("obj-1", "Объект 1 (раб)", "ctr-1", "Договор 1", 1, 1),
        _object_mapping_row("obj-missing", "Нет объекта", "ctr-2", "Договор 2", 1, 1),
    ]
    content = _build_workbook(rows)

    apply_resp = await client.post(
        f"{IMPORT_URL}/apply",
        files={"file": ("catalogs.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert apply_resp.status_code == 422

    assert db_session.execute(select(Contract).where(Contract.code == "ctr-1")).scalar_one_or_none() is None
    assert db_session.execute(select(ObjectContract)).scalar_one_or_none() is None


async def test_import_object_mappings_duplicate_pair_rejected(client, db_session):
    rows = _build_object_only()
    rows["ObjectMappings"] = [
        _object_mapping_row("obj-1", "Объект 1 (раб)", "ctr-1", "Договор 1", 1, 1),
        _object_mapping_row("obj-1", "Объект 1 (раб)", "ctr-1", "Договор 1 дубль", 1, 1),
    ]
    content = _build_workbook(rows)

    validate_resp = await client.post(
        f"{IMPORT_URL}/validate",
        files={"file": ("catalogs.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert validate_resp.status_code == 200
    body = validate_resp.json()
    assert body["valid"] is False
    assert any(
        "Дублирующаяся пара (object_code, contract_code)" in err["message"]
        for err in body["errors"]
    )

    apply_resp = await client.post(
        f"{IMPORT_URL}/apply",
        files={"file": ("catalogs.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert apply_resp.status_code == 422


async def test_simple_object_workbook_is_converted_and_idempotent(client, db_session):
    db_session.add_all(
        [
            Stage(code="reconstruction", name="Реконструкция", active=True),
            Stage(code="cable-04", name="КЛ 0,4", active=True),
            Stage(code="inactive", name="Неактивный", active=False),
        ]
    )
    db_session.commit()
    content = _build_simple_objects_workbook()

    validate_resp = await client.post(
        f"{IMPORT_URL}/validate",
        files={"file": ("objects.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert validate_resp.status_code == 200
    validation = validate_resp.json()
    assert validation["valid"] is True
    assert validation["preview"]["Objects"]["create"] == 3

    for _ in range(2):
        apply_resp = await client.post(
            f"{IMPORT_URL}/apply",
            files={"file": ("objects.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert apply_resp.status_code == 200

    objects = db_session.execute(select(Object).order_by(Object.name)).scalars().all()
    contracts = db_session.execute(select(Contract)).scalars().all()
    links = db_session.execute(select(ObjectContract)).scalars().all()
    object_stages = db_session.execute(select(ObjectStage)).scalars().all()
    assert [obj.name for obj in objects] == ["Обводный канал", "Пискаревский", "Тамбасова"]
    assert len(contracts) == 3
    assert len(links) == 3
    assert len(object_stages) == 6
    active_stage_ids = set(
        db_session.execute(select(Stage.id).where(Stage.active)).scalars().all()
    )
    assert {link.stage_id for link in object_stages} == active_stage_ids
    tambasova = next(obj for obj in objects if obj.name == "Тамбасова")
    assert sum(link.object_id == tambasova.id for link in links) == 2


async def test_simple_object_workbook_requires_an_active_stage(client, db_session):
    content = _build_simple_objects_workbook()

    apply_resp = await client.post(
        f"{IMPORT_URL}/apply",
        files={
            "file": (
                "objects.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert apply_resp.status_code == 422
    assert "нет активных этапов" in str(apply_resp.json()["detail"])
    assert db_session.execute(select(Object)).scalar_one_or_none() is None


async def test_unknown_workbook_is_rejected(client):
    wb = Workbook()
    wb.active.title = "Неизвестный лист"
    wb.active.append(["другая", "структура"])
    buffer = BytesIO()
    wb.save(buffer)

    validate_resp = await client.post(
        f"{IMPORT_URL}/validate",
        files={"file": ("unknown.xlsx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert validate_resp.status_code == 200
    assert validate_resp.json()["valid"] is False

    apply_resp = await client.post(
        f"{IMPORT_URL}/apply",
        files={"file": ("unknown.xlsx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert apply_resp.status_code == 422


async def test_unimplemented_staged_import_route_is_not_published(client):
    response = await client.post(f"{IMPORT_URL}/123/apply")
    openapi = (await client.get("/openapi.json")).json()

    assert response.status_code in {404, 405}
    assert "/api/catalogs/import/{import_id}/apply" not in openapi["paths"]
