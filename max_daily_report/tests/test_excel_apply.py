from __future__ import annotations

from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.main import app
from app.models.catalogs import (
    Contractor,
    EquipmentType,
    Object,
    ObjectStage,
    Stage,
    Unit,
    WorkMethod,
    WorkType,
    WorkTypeMethod,
)
from app.services.excel_service import SHEET_COLUMNS, SHEET_ORDER, build_template

client = TestClient(app)


import pytest  # noqa: E402


@pytest.fixture(autouse=True)
def _admin_dev_user(db_session):
    """Import endpoints require manager/admin — pre-create dev-user as admin."""
    from app.models.users import User

    db_session.add(User(max_user_id="dev-user", full_name="Dev User", role="admin"))
    db_session.commit()


def _build_workbook(rows_by_sheet: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for sheet_name in SHEET_ORDER:
        ws = wb.create_sheet(title=sheet_name)
        ws.append(SHEET_COLUMNS[sheet_name])
        for row in rows_by_sheet.get(sheet_name, []):
            ws.append([row.get(col) for col in SHEET_COLUMNS[sheet_name]])
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_apply_creates_catalogs(db_session):
    rows = {
        "Contractors": [
            {"code": "c1", "name": "Подрядчик 1", "active": 1, "sort_order": 1},
        ],
        "Units": [
            {"code": "m3", "name": "кубометр", "symbol": "м³", "active": 1, "sort_order": 1},
        ],
        "Stages": [
            {"code": "stg1", "name": "Этап 1", "active": 1, "sort_order": 1},
        ],
        "Objects": [
            {
                "code": "obj1",
                "name": "Объект 1",
                "execution_method": "own",
                "default_contractor_code": "c1",
                "active": 1,
                "sort_order": 1,
            },
        ],
        "ObjectStages": [
            {"object_code": "obj1", "stage_code": "stg1", "active": 1},
        ],
        "Equipment": [
            {"code": "eq1", "name": "Экскаватор", "default_unit_code": "m3", "active": 1, "sort_order": 1},
        ],
        "WorkTypes": [
            {"code": "wt1", "name": "Работа 1", "default_unit_code": "m3", "active": 1, "sort_order": 1},
        ],
        "WorkMethods": [
            {"code": "wm1", "name": "Способ 1", "active": 1, "sort_order": 1},
        ],
        "WorkTypeMethods": [
            {"work_type_code": "wt1", "work_method_code": "wm1", "active": 1},
        ],
    }
    response = client.post(
        "/api/catalogs/import/apply",
        files={"file": ("catalogs.xlsx", _build_workbook(rows), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "applied"


def test_apply_creates_rows_in_db(db_session):
    rows = {
        "Contractors": [
            {"code": "c2", "name": "Подрядчик 2", "active": 1, "sort_order": 1},
        ],
        "Units": [
            {"code": "m2", "name": "метр", "symbol": "м", "active": 1, "sort_order": 1},
        ],
        "Stages": [
            {"code": "stg2", "name": "Этап 2", "active": 1, "sort_order": 1},
        ],
        "Objects": [
            {
                "code": "obj2",
                "name": "Объект 2",
                "execution_method": "contractor",
                "default_contractor_code": "c2",
                "active": 1,
                "sort_order": 1,
            },
        ],
        "ObjectStages": [
            {"object_code": "obj2", "stage_code": "stg2", "active": 1},
        ],
        "Equipment": [
            {"code": "eq2", "name": "Бульдозер", "default_unit_code": "m2", "active": 1, "sort_order": 1},
        ],
        "WorkTypes": [
            {"code": "wt2", "name": "Работа 2", "default_unit_code": "m2", "active": 1, "sort_order": 1},
        ],
        "WorkMethods": [
            {"code": "wm2", "name": "Способ 2", "active": 1, "sort_order": 1},
        ],
        "WorkTypeMethods": [
            {"work_type_code": "wt2", "work_method_code": "wm2", "active": 1},
        ],
    }
    response = client.post(
        "/api/catalogs/import/apply",
        files={"file": ("catalogs.xlsx", _build_workbook(rows), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200

    result = db_session.execute(select(Contractor).where(Contractor.code == "c2"))
    c = result.scalar_one()
    assert c.name == "Подрядчик 2"
    assert db_session.execute(select(Unit)).scalar_one_or_none() is not None
    assert db_session.execute(select(Stage)).scalar_one_or_none() is not None
    assert db_session.execute(select(Object)).scalar_one_or_none() is not None
    assert db_session.execute(select(ObjectStage)).scalar_one_or_none() is not None
    assert db_session.execute(select(EquipmentType)).scalar_one_or_none() is not None
    assert db_session.execute(select(WorkType)).scalar_one_or_none() is not None
    assert db_session.execute(select(WorkMethod)).scalar_one_or_none() is not None
    assert db_session.execute(select(WorkTypeMethod)).scalar_one_or_none() is not None


def test_apply_rollback_on_invalid_reference(db_session):
    rows = {
        "Objects": [
            {
                "code": "obj-bad",
                "name": "Объект без подрядчика",
                "execution_method": "own",
                "default_contractor_code": "missing-contractor",
                "active": 1,
                "sort_order": 1,
            },
        ],
    }
    response = client.post(
        "/api/catalogs/import/apply",
        files={"file": ("bad.xlsx", _build_workbook(rows), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 422


def test_apply_rollback_on_invalid_reference_db(db_session):
    rows = {
        "Contractors": [
            {"code": "c-rollback", "name": "Кандидат", "active": 1, "sort_order": 1},
        ],
        "Objects": [
            {
                "code": "obj-bad",
                "name": "Объект без подрядчика",
                "execution_method": "own",
                "default_contractor_code": "missing-contractor",
                "active": 1,
                "sort_order": 1,
            },
        ],
    }
    response = client.post(
        "/api/catalogs/import/apply",
        files={"file": ("bad.xlsx", _build_workbook(rows), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 422
    result = db_session.execute(
        select(Contractor).where(Contractor.code == "c-rollback")
    )
    assert result.scalar_one_or_none() is None


def test_apply_rejects_validation_errors(db_session):
    rows = {
        "Objects": [
            {
                "code": "obj-bad",
                "name": "Объект",
                "execution_method": "invalid",
                "default_contractor_code": None,
                "active": 1,
                "sort_order": 1,
            },
        ],
    }
    response = client.post(
        "/api/catalogs/import/apply",
        files={"file": ("bad.xlsx", _build_workbook(rows), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 422


def test_export_template_roundtrip(db_session):
    buffer = build_template()
    response = client.post(
        "/api/catalogs/import/validate",
        files={"file": ("template.xlsx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200
    body = response.json()
    assert body["valid"] is True


def test_export_catalogs_endpoint(db_session):
    response = client.get("/api/catalogs/export.xlsx")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    wb = load_workbook(filename=BytesIO(response.content))
    assert "Objects" in wb.sheetnames
    assert "Contractors" in wb.sheetnames
