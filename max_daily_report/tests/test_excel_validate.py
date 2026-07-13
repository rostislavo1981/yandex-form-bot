from __future__ import annotations

from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.main import app
from app.services.excel_service import SHEET_COLUMNS, SHEET_ORDER, build_template

client = TestClient(app)


def _build_workbook(rows_by_sheet: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for sheet_name in SHEET_ORDER:
        ws = wb.create_sheet(title=sheet_name)
        ws.append(SHEET_COLUMNS[sheet_name])
        for row in rows_by_sheet.get(sheet_name, []):
            ws.append([row.get(col) for col in SHEET_COLUMNS[sheet_name]])
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_validate_valid_preview() -> None:
    rows = {
        "Contractors": [
            {"code": "own", "name": "Собственные", "active": 1, "sort_order": 1},
            {"code": "sub-1", "name": "Подрядчик 1", "active": 1, "sort_order": 2},
        ],
        "Units": [
            {"code": "m3", "name": "кубометр", "symbol": "м³", "active": 1, "sort_order": 1},
        ],
        "Objects": [
            {
                "code": "obj-1",
                "name": "Объект 1",
                "execution_method": "own",
                "default_contractor_code": "own",
                "active": 1,
                "sort_order": 1,
            },
        ],
    }
    response = client.post(
        "/api/catalogs/import/validate",
        files={"file": ("catalogs.xlsx", _build_workbook(rows), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200
    body = response.json()
    assert body["valid"] is True
    assert body["preview"]["Contractors"]["create"] == 2


def test_validate_rejects_duplicate_code() -> None:
    rows = {
        "Contractors": [
            {"code": "dup", "name": "Один", "active": 1, "sort_order": 1},
            {"code": "dup", "name": "Два", "active": 1, "sort_order": 2},
        ],
    }
    response = client.post(
        "/api/catalogs/import/validate",
        files={"file": ("catalogs.xlsx", _build_workbook(rows), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200
    body = response.json()
    assert body["valid"] is False
    assert any("Дублирующийся code" in err["message"] for err in body["errors"])


def test_validate_rejects_invalid_enum() -> None:
    rows = {
        "Objects": [
            {
                "code": "obj-bad",
                "name": "Объект",
                "execution_method": "invalid",
                "default_contractor_code": None,
                "active": 1,
                "sort_order": 1,
            },
        ],
    }
    response = client.post(
        "/api/catalogs/import/validate",
        files={"file": ("catalogs.xlsx", _build_workbook(rows), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200
    body = response.json()
    assert body["valid"] is False
    assert any("execution_method" in err["message"] for err in body["errors"])


def test_validate_rejects_non_xlsx() -> None:
    response = client.post(
        "/api/catalogs/import/validate",
        files={"file": ("catalogs.txt", b"not excel", "text/plain")},
    )
    assert response.status_code == 400


def test_build_template_contains_all_sheets() -> None:
    buffer = build_template()
    from openpyxl import load_workbook

    loaded = load_workbook(filename=BytesIO(buffer.getvalue()))
    for sheet_name in SHEET_ORDER:
        assert sheet_name in loaded.sheetnames
