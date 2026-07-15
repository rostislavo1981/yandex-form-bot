from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import settings
from app.main import create_app
from app.models.catalogs import Contractor, EquipmentType, Object, Stage, Unit, WorkType
from app.models.contracts import Contract
from app.models.reports import (
    DailyReport,
    ReportEquipment,
    ReportWork,
)
from app.models.users import MAXGroup, User


@pytest.fixture
def client() -> TestClient:
    return TestClient(create_app())


@pytest.fixture(autouse=True)
def _dev_env(monkeypatch):
    monkeypatch.setattr(settings, "app_env", "dev")
    monkeypatch.setattr(settings, "debug", True)


def _seed(db_session):
    dev_user = User(max_user_id="dev-user", full_name="Dev User", role="manager")
    db_session.add(dev_user)
    user = User(max_user_id="max-u1", full_name="Иван", role="responsible")
    obj_a = Object(code="OBJ-A", name="Объект A", active=True)
    obj_b = Object(code="OBJ-B", name="Объект B", active=True)
    stage = Stage(code="STG", name="Этап", active=True)
    unit_m3 = Unit(code="m3", name="м³", symbol="м³", active=True)
    unit_h = Unit(code="h", name="час", symbol="ч", active=True)
    unit_t = Unit(code="t", name="тонна", symbol="т", active=True)
    eq_type = EquipmentType(code="EXC", name="Экскаватор", active=True)
    work_type = WorkType(code="DIG", name="Земляные работы", active=True)
    group = MAXGroup(chat_id="chat-1", title="Группа", active=True)
    contractor = Contractor(code="CON", name="Подрядчик", active=True)
    contract1 = Contract(code="C1", full_name="Договор Основной", active=True)
    contract2 = Contract(code="C2", full_name="Договор Дополнительный", active=True)
    db_session.add_all([
        user, obj_a, obj_b, stage, unit_m3, unit_h, unit_t,
        eq_type, work_type, group, contractor, contract1, contract2,
    ])
    db_session.flush()
    return {
        "user": user,
        "obj_a": obj_a,
        "obj_b": obj_b,
        "stage": stage,
        "unit_m3": unit_m3,
        "unit_h": unit_h,
        "unit_t": unit_t,
        "eq_type": eq_type,
        "work_type": work_type,
        "contractor": contractor,
        "contract1": contract1,
        "contract2": contract2,
    }


def _make_report(
    db_session,
    report_date: date,
    user: User,
    obj: Object,
    stage: Stage,
    eq_type: EquipmentType,
    work_type: WorkType,
    contractor: Contractor,
    contract: Contract,
    unit_h: Unit,
    unit_m3: Unit,
    work_quantity: Decimal = Decimal("10.00"),
    eq_quantity: Decimal = Decimal("2.00"),
) -> DailyReport:
    report = DailyReport(
        report_date=report_date,
        responsible_user_id=user.id,
        object_id=obj.id,
        stage_id=stage.id,
        contractor_id=contractor.id,
        contract_id=contract.id,
        object_name_snapshot=obj.name,
        contract_code_snapshot=contract.code,
        contract_full_name_snapshot=contract.full_name,
        staff_itr=1,
        staff_internal=1,
        staff_external=1,
        status="submitted",
        idempotency_key=f"key-{obj.code}-{report_date.isoformat()}-{contract.code}",
    )
    db_session.add(report)
    db_session.flush()
    db_session.add(
        ReportEquipment(
            report_id=report.id,
            equipment_type_id=eq_type.id,
            equipment_name_snapshot="Экскаватор",
            ownership="own",
            unit_id=unit_h.id,
            unit_name_snapshot=unit_h.symbol,
            quantity=eq_quantity,
        )
    )
    db_session.add(
        ReportWork(
            report_id=report.id,
            work_type_id=work_type.id,
            work_name_snapshot="Земляные работы",
            unit_id=unit_m3.id,
            unit_name_snapshot=unit_m3.symbol,
            quantity=work_quantity,
        )
    )
    return report


def test_timesheet_separated_by_object(client: TestClient, db_session):
    s = _seed(db_session)
    _make_report(
        db_session, date(2026, 7, 14), s["user"], s["obj_a"], s["stage"],
        s["eq_type"], s["work_type"], s["contractor"], s["contract1"], s["unit_h"], s["unit_m3"],
    )
    _make_report(
        db_session, date(2026, 7, 14), s["user"], s["obj_b"], s["stage"],
        s["eq_type"], s["work_type"], s["contractor"], s["contract1"], s["unit_h"], s["unit_m3"],
    )
    db_session.commit()

    response = client.get(
        f"/api/timesheet/{s['obj_a'].id}?date_from=2026-07-14&date_to=2026-07-14"
    )
    assert response.status_code == 200
    data = response.json()
    assert data["object_code"] == "OBJ-A"
    assert data["object_name"] == "Объект A"
    assert "raw_reports" in data
    for row in data["raw_reports"]:
        assert row["object_code"] == "OBJ-A"
    assert len(data["raw_reports"]) > 0


def test_timesheet_excel_contains_contract_info(client: TestClient, db_session):
    s = _seed(db_session)
    _make_report(
        db_session, date(2026, 7, 14), s["user"], s["obj_a"], s["stage"],
        s["eq_type"], s["work_type"], s["contractor"], s["contract1"], s["unit_h"], s["unit_m3"],
        work_quantity=Decimal("10.00"),
    )
    _make_report(
        db_session, date(2026, 7, 14), s["user"], s["obj_a"], s["stage"],
        s["eq_type"], s["work_type"], s["contractor"], s["contract2"], s["unit_h"], s["unit_m3"],
        work_quantity=Decimal("20.00"),
    )
    db_session.commit()

    response = client.get(
        f"/api/timesheet/{s['obj_a'].id}/export.xlsx?date_from=2026-07-14&date_to=2026-07-14"
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    wb = load_workbook(BytesIO(response.content))
    assert "Исходные отчёты" in wb.sheetnames
    ws = wb["Исходные отчёты"]
    header = [c.value for c in ws[1]]
    assert "Договор код" in header
    assert "Договор название" in header

    code_idx = header.index("Договор код")
    name_idx = header.index("Договор название")
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert rows
    codes = {r[code_idx] for r in rows}
    names = {r[name_idx] for r in rows}
    assert "C1" in codes
    assert "C2" in codes
    assert "Договор Основной" in names
    assert "Договор Дополнительный" in names


def test_timesheet_sums_do_not_mix_units(client: TestClient, db_session):
    s = _seed(db_session)
    report = _make_report(
        db_session, date(2026, 7, 14), s["user"], s["obj_a"], s["stage"],
        s["eq_type"], s["work_type"], s["contractor"], s["contract1"], s["unit_h"], s["unit_m3"],
    )
    # same work name, different units — must be summed separately
    db_session.add(
        ReportWork(
            report_id=report.id,
            work_type_id=s["work_type"].id,
            work_name_snapshot="Земляные работы",
            unit_id=s["unit_t"].id,
            unit_name_snapshot=s["unit_t"].symbol,
            quantity=Decimal("5.00"),
        )
    )
    db_session.commit()

    response = client.get(
        f"/api/timesheet/{s['obj_a'].id}?date_from=2026-07-14&date_to=2026-07-14"
    )
    assert response.status_code == 200
    data = response.json()

    work_rows = [r for r in data["rows"] if r["category"] == "work"]
    assert len(work_rows) == 2
    units = {r["unit"] for r in work_rows}
    assert units == {s["unit_m3"].symbol, s["unit_t"].symbol}

    raw_work = [r for r in data["raw_reports"] if r["category"] == "work"]
    assert len(raw_work) == 2
    raw_units = {r["unit"] for r in raw_work}
    assert raw_units == {s["unit_m3"].symbol, s["unit_t"].symbol}


def test_two_reports_same_object_different_contracts_in_details(client: TestClient, db_session):
    s = _seed(db_session)
    _make_report(
        db_session, date(2026, 7, 14), s["user"], s["obj_a"], s["stage"],
        s["eq_type"], s["work_type"], s["contractor"], s["contract1"], s["unit_h"], s["unit_m3"],
    )
    _make_report(
        db_session, date(2026, 7, 14), s["user"], s["obj_a"], s["stage"],
        s["eq_type"], s["work_type"], s["contractor"], s["contract2"], s["unit_h"], s["unit_m3"],
    )
    db_session.commit()

    response = client.get(
        f"/api/timesheet/{s['obj_a'].id}?date_from=2026-07-14&date_to=2026-07-14"
    )
    assert response.status_code == 200
    data = response.json()

    codes = {r["contract_code"] for r in data["raw_reports"]}
    names = {r["contract_full_name"] for r in data["raw_reports"]}
    assert "C1" in codes
    assert "C2" in codes
    assert "Договор Основной" in names
    assert "Договор Дополнительный" in names
