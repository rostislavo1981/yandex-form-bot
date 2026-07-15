from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import settings
from app.main import create_app
from app.models.catalogs import EquipmentType, Object, Stage, Unit, WorkType
from app.models.reports import (
    DailyReport,
    ReportEquipment,
    ReportObligation,
    ReportWork,
    ResponsibleObjectAssignment,
)
from app.models.users import MAXGroup, User


@pytest.fixture
def client() -> TestClient:
    return TestClient(create_app())


@pytest.fixture(autouse=True)
def _dev_env(monkeypatch):
    monkeypatch.setattr(settings, "app_env", "dev")
    monkeypatch.setattr(settings, "debug", True)


def _seed_excel_data(db_session):
    dev_user = User(max_user_id="dev-user", full_name="Dev User", role="manager")
    db_session.add(dev_user)
    user = User(max_user_id="max-u1", full_name="Иван", role="responsible")
    obj_a = Object(code="OBJ-A", name="Объект A", active=True)
    obj_b = Object(code="OBJ-B", name="Объект B", active=True)
    stage = Stage(code="STG", name="Этап", active=True)
    unit_m3 = Unit(code="m3", name="м³", symbol="м³", active=True)
    unit_h = Unit(code="h", name="час", symbol="ч", active=True)
    eq_type = EquipmentType(code="EXC", name="Экскаватор", active=True)
    work_type = WorkType(code="DIG", name="Земляные работы", active=True)
    group = MAXGroup(chat_id="chat-1", title="Группа", active=True)
    db_session.add_all([user, obj_a, obj_b, stage, unit_m3, unit_h, eq_type, work_type, group])
    db_session.flush()

    assignment_a = ResponsibleObjectAssignment(
        user_id=user.id,
        object_id=obj_a.id,
        active_from=date(2026, 7, 1),
        active_to=date(2026, 7, 31),
        schedule_type="daily",
    )
    db_session.add(assignment_a)
    db_session.flush()
    for d in [date(2026, 7, 14), date(2026, 7, 15)]:
        db_session.add(
            ReportObligation(
                report_date=d,
                assignment_id=assignment_a.id,
                user_id=user.id,
                object_id=obj_a.id,
                status="pending",
            )
        )
    db_session.flush()

    report = DailyReport(
        report_date=date(2026, 7, 14),
        responsible_user_id=user.id,
        object_id=obj_a.id,
        stage_id=stage.id,
        object_name_snapshot="Объект A",
        staff_itr=1,
        staff_internal=2,
        staff_external=1,
        status="submitted",
        idempotency_key="key-1",
    )
    db_session.add(report)
    db_session.flush()
    db_session.add(
        ReportEquipment(
            report_id=report.id,
            equipment_type_id=eq_type.id,
            equipment_name_snapshot="Экскаватор",
            ownership="own",
            unit_id=unit_h.id,
            unit_name_snapshot="ч",
            quantity=Decimal("8.00"),
        )
    )
    db_session.add(
        ReportWork(
            report_id=report.id,
            work_type_id=work_type.id,
            work_name_snapshot="Земляные работы",
            unit_id=unit_m3.id,
            unit_name_snapshot="м³",
            quantity=Decimal("130.50"),
        )
    )
    db_session.commit()
    return obj_a, obj_b


def test_export_timesheet_workbook(client: TestClient, db_session):
    obj_a, obj_b = _seed_excel_data(db_session)
    response = client.get(
        f"/api/timesheet/{obj_a.id}/export.xlsx?date_from=2026-07-14&date_to=2026-07-15"
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    wb = load_workbook(BytesIO(response.content))
    assert "Общая сводка" in wb.sheetnames
    assert "Статус отправки" in wb.sheetnames
    assert "OBJ-A" in wb.sheetnames
    assert "Исходные отчёты" in wb.sheetnames

    object_ws = wb["OBJ-A"]
    # header contains days
    assert object_ws.cell(row=1, column=4).value == "2026-07-14"
    assert object_ws.cell(row=1, column=5).value == "2026-07-15"

    # find work row
    work_row = None
    for row in object_ws.iter_rows(min_row=2, values_only=True):
        if row[0] == "work":
            work_row = row
            break
    assert work_row is not None
    assert work_row[3] == 130.5 or str(work_row[3]) == "130.5"


def test_export_timesheet_object_b_not_affected(client: TestClient, db_session):
    obj_a, obj_b = _seed_excel_data(db_session)
    response = client.get(
        f"/api/timesheet/{obj_a.id}/export.xlsx?date_from=2026-07-14&date_to=2026-07-15"
    )
    wb = load_workbook(BytesIO(response.content))
    assert "OBJ-B" not in wb.sheetnames
