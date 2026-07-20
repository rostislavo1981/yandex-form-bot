from __future__ import annotations

import hashlib
from io import BytesIO
from typing import Any

import sqlalchemy as sa
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.catalogs import (
    Contractor,
    EquipmentType,
    Object,
    ObjectStage,
    Stage,
    Unit,
    WorkMethod,
    WorkType,
    WorkTypeMethod,
)
from app.models.contracts import Contract, ObjectContract
from app.models.operations import CatalogImport
from app.models.reports import ResponsibleObjectAssignment
from app.models.users import User
from app.repos.catalogs import CatalogRepo

_EMPTY_CONTRACT_MARKERS = ("", "??", "нет пока договора")
_SIMPLE_OBJECTS_HEADER = "краткое название"
_CATALOG_IMPORT_LOCK_ID = 6146747792393678674

SHEET_COLUMNS = {
    "Users": ["code", "name", "role", "active", "sort_order"],
    "Objects": [
        "code",
        "name",
        "short_title",
        "full_title",
        "execution_method",
        "default_contractor_code",
        "active",
        "sort_order",
    ],
    "Stages": ["code", "name", "active", "sort_order"],
    "ObjectStages": ["object_code", "stage_code", "active"],
    "Contractors": ["code", "name", "active", "sort_order"],
    "Units": ["code", "name", "symbol", "active", "sort_order"],
    "Equipment": ["code", "name", "default_unit_code", "active", "sort_order"],
    "WorkTypes": ["code", "name", "default_unit_code", "active", "sort_order"],
    "WorkMethods": ["code", "name", "active", "sort_order"],
    "WorkTypeMethods": ["work_type_code", "work_method_code", "active"],
    "Assignments": [
        "user_code",
        "object_code",
        "active_from",
        "active_to",
        "schedule_type",
        "active",
    ],
    "ObjectMappings": [
        "object_code",
        "short_name",
        "contract_code",
        "full_name",
        "primary",
        "active",
    ],
}

SHEET_ORDER = [
    "Users",
    "Contractors",
    "Units",
    "Objects",
    "Stages",
    "ObjectStages",
    "Equipment",
    "WorkTypes",
    "WorkMethods",
    "WorkTypeMethods",
    "Assignments",
    "ObjectMappings",
]


def _normalize_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _normalize_bool(value: Any) -> bool:
    if value is None:
        return True
    text = str(value).strip().lower()
    if text in ("1", "true", "yes", "да"):
        return True
    if text in ("0", "false", "no", "нет"):
        return False
    return bool(text)


def _stable_import_code(prefix: str, value: str) -> str:
    """Build a deterministic short code for simple two-column workbooks."""
    normalized = " ".join(value.lower().split())
    digest = hashlib.sha256(normalized.encode("utf-8")).hexdigest()[:12].upper()
    return f"{prefix}-{digest}"


def _convert_simple_object_workbook(source: Workbook) -> Workbook | None:
    """Convert the user's simple ``short name | contract(s)`` workbook.

    The Yandex Form source has a title above the table, the short name in one
    column and one or more official names in the columns to its right.  It is
    intentionally accepted directly so an administrator does not have to
    manually rebuild the full multi-sheet catalog template.
    """
    source_sheet = None
    header_row = None
    short_col = None
    for ws in source.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if _normalize_text(cell.value).lower() == _SIMPLE_OBJECTS_HEADER:
                    source_sheet = ws
                    header_row = cell.row
                    short_col = cell.column
                    break
            if source_sheet is not None:
                break
        if source_sheet is not None:
            break

    if source_sheet is None or header_row is None or short_col is None:
        return None

    converted = Workbook()
    converted.remove(converted.active)
    objects_ws = converted.create_sheet("Objects")
    mappings_ws = converted.create_sheet("ObjectMappings")
    objects_ws.append(SHEET_COLUMNS["Objects"])
    mappings_ws.append(SHEET_COLUMNS["ObjectMappings"])

    sort_order = 0
    for row_idx in range(header_row + 1, source_sheet.max_row + 1):
        short_name = _normalize_text(source_sheet.cell(row_idx, short_col).value)
        if not short_name:
            continue
        sort_order += 1
        object_code = _stable_import_code("OBJ", short_name)

        contract_values = [
            _normalize_text(source_sheet.cell(row_idx, col_idx).value)
            for col_idx in range(short_col + 1, source_sheet.max_column + 1)
        ]
        contract_values = [value for value in contract_values if value]
        if not contract_values:
            contract_values = [""]

        primary_full_name = ""
        for contract_idx, full_name in enumerate(contract_values):
            marker = full_name.lower()
            if marker in _EMPTY_CONTRACT_MARKERS:
                contract_code = marker
                normalized_full_name = ""
            else:
                contract_code = _stable_import_code("CTR", full_name)
                normalized_full_name = full_name
                if contract_idx == 0:
                    primary_full_name = normalized_full_name
            mappings_ws.append(
                [
                    object_code,
                    short_name,
                    contract_code,
                    normalized_full_name,
                    int(contract_idx == 0),
                    1,
                ]
            )

        objects_ws.append(
            [
                object_code,
                short_name,
                short_name,
                primary_full_name,
                "own",
                "",
                1,
                sort_order,
            ]
        )
    return converted


def build_template() -> BytesIO:
    """Build an empty catalog import template workbook."""
    wb = Workbook()
    wb.remove(wb.active)
    for sheet_name in SHEET_ORDER:
        ws = wb.create_sheet(title=sheet_name)
        columns = SHEET_COLUMNS[sheet_name]
        ws.append(columns)
        for col_idx, _col in enumerate(columns, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 20
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _read_sheet_rows(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [_normalize_text(cell).lower() for cell in rows[0]]
    result = []
    for row in rows[1:]:
        record = {}
        for idx, header in enumerate(headers):
            record[header] = row[idx] if idx < len(row) else None
        result.append(record)
    return result


class CatalogImportValidator:
    """Validate an uploaded catalog Excel workbook and produce a preview."""

    def __init__(self, workbook_bytes: bytes) -> None:
        self._wb = load_workbook(filename=BytesIO(workbook_bytes), data_only=True)
        self._simple_object_workbook = False
        if not set(self._wb.sheetnames).intersection(SHEET_ORDER):
            converted = _convert_simple_object_workbook(self._wb)
            if converted is not None:
                self._wb = converted
                self._simple_object_workbook = True
        self._errors: list[dict] = []
        self._warnings: list[dict] = []
        self._preview: dict = {}
        self._seen_object_mapping_pairs: dict[tuple[str, str], int] = {}

    def validate(self) -> dict:
        self._errors = []
        self._warnings = []
        self._preview = {}
        self._seen_object_mapping_pairs = {}
        recognized_sheets = set(self._wb.sheetnames).intersection(SHEET_ORDER)
        if not recognized_sheets:
            self._add_error(
                "Workbook",
                0,
                "Не найден ни один поддерживаемый лист или таблица с заголовком "
                "'краткое название'",
            )
        for sheet_name in SHEET_ORDER:
            ws = self._wb[sheet_name] if sheet_name in self._wb.sheetnames else None
            if ws is None:
                continue
            rows = _read_sheet_rows(ws)
            self._preview[sheet_name] = {"create": 0, "update": 0, "deactivate": 0}
            self._validate_sheet(sheet_name, rows)
        return {
            "valid": not self._errors,
            "errors": self._errors,
            "warnings": self._warnings,
            "preview": self._preview,
        }

    @property
    def preview(self) -> dict:
        return self._preview

    @property
    def errors(self) -> list[dict]:
        return self._errors

    @property
    def is_simple_object_workbook(self) -> bool:
        return self._simple_object_workbook

    def iter_rows(self, sheet_name: str) -> list[dict]:
        ws = self._wb[sheet_name] if sheet_name in self._wb.sheetnames else None
        if ws is None:
            return []
        return _read_sheet_rows(ws)

    def _add_error(self, sheet: str, row: int, message: str) -> None:
        self._errors.append({"sheet": sheet, "row": row, "message": message})

    def _add_warning(self, sheet: str, row: int, message: str) -> None:
        self._warnings.append(
            {"sheet": sheet, "row": row, "message": message, "level": "warning"}
        )

    def _validate_sheet(self, sheet_name: str, rows: list[dict]) -> None:
        required = {"code"} if "code" in SHEET_COLUMNS[sheet_name] else set()
        codes_seen: dict[str, int] = {}

        for idx, row in enumerate(rows, start=2):
            for col in required:
                value = _normalize_text(row.get(col))
                if not value:
                    self._add_error(sheet_name, idx, f"Пустое обязательное поле '{col}'")
                    continue

            code = _normalize_text(row.get("code"))
            if code:
                if code in codes_seen:
                    self._add_error(
                        sheet_name,
                        idx,
                        f"Дублирующийся code '{code}' (строка {codes_seen[code]})",
                    )
                else:
                    codes_seen[code] = idx

            if sheet_name == "Objects":
                method = _normalize_text(row.get("execution_method"))
                if method and method not in ("own", "contractor"):
                    self._add_error(
                        sheet_name, idx, f"execution_method '{method}' не из own|contractor"
                    )

            if sheet_name == "ObjectStages":
                obj_code = _normalize_text(row.get("object_code"))
                stage_code = _normalize_text(row.get("stage_code"))
                if not obj_code:
                    self._add_error(sheet_name, idx, "Пустое обязательное поле 'object_code'")
                if not stage_code:
                    self._add_error(sheet_name, idx, "Пустое обязательное поле 'stage_code'")

            if sheet_name == "Users":
                role = _normalize_text(row.get("role"))
                if role and role not in ("responsible", "manager", "admin"):
                    self._add_error(sheet_name, idx, f"role '{role}' не из responsible|manager|admin")

            if sheet_name == "ObjectMappings":
                self._validate_object_mappings(idx, row)

            active = _normalize_bool(row.get("active"))
            if active:
                self._preview[sheet_name]["create" if code else "update"] += 1
            else:
                self._preview[sheet_name]["deactivate"] += 1

    def _validate_object_mappings(self, idx: int, row: dict) -> None:
        object_code = _normalize_text(row.get("object_code"))
        short_name = _normalize_text(row.get("short_name"))
        contract_code = _normalize_text(row.get("contract_code"))
        full_name = _normalize_text(row.get("full_name"))

        if not object_code:
            self._add_error("ObjectMappings", idx, "Пустое обязательное поле 'object_code'")
        if not short_name:
            self._add_error("ObjectMappings", idx, "Пустое обязательное поле 'short_name'")

        if contract_code in _EMPTY_CONTRACT_MARKERS:
            self._add_warning(
                "ObjectMappings",
                idx,
                "contract_code указан как отсутствующий — договор не будет создан",
            )
        elif contract_code and not full_name:
            self._add_error(
                "ObjectMappings",
                idx,
                "Пустое обязательное поле 'full_name' при указанном contract_code",
            )

        if object_code and contract_code and contract_code not in _EMPTY_CONTRACT_MARKERS:
            pair = (object_code, contract_code)
            if pair in self._seen_object_mapping_pairs:
                self._add_error(
                    "ObjectMappings",
                    idx,
                    f"Дублирующаяся пара (object_code, contract_code) "
                    f"'{object_code}', '{contract_code}' (строка {self._seen_object_mapping_pairs[pair]})",
                )
            else:
                self._seen_object_mapping_pairs[pair] = idx


class CatalogImportApplier:
    """Apply a validated catalog workbook inside one transaction."""

    def __init__(self, session: AsyncSession) -> None:
        self._session = session
        self._repo = CatalogRepo(session)
        self._contractors: dict[str, Contractor] = {}
        self._units: dict[str, Unit] = {}
        self._stages: dict[str, Stage] = {}
        self._objects: dict[str, Object] = {}
        self._methods: dict[str, WorkMethod] = {}
        self._work_types: dict[str, WorkType] = {}

    async def apply(self, validator: CatalogImportValidator) -> CatalogImport:
        # A workbook touches several related tables. Serialize apply operations so
        # two simultaneous admin clicks cannot both decide that the same code is
        # absent and then race on a unique constraint.
        await self._session.execute(
            sa.text("SELECT pg_advisory_xact_lock(:lock_id)"),
            {"lock_id": _CATALOG_IMPORT_LOCK_ID},
        )
        import_record = CatalogImport(
            filename="applied.xlsx",
            status="validated",
            summary_json=validator.preview,
            errors_json=validator.errors,
        )
        self._session.add(import_record)
        await self._session.flush()

        if not validator.validate()["valid"]:
            import_record.status = "failed"
            await self._session.commit()
            raise ValueError("Импорт содержит ошибки валидации")

        await self._load_existing()
        await self._apply_users(validator)
        await self._apply_contractors(validator)
        await self._apply_units(validator)
        await self._apply_stages(validator)
        await self._apply_objects(validator)
        if validator.is_simple_object_workbook:
            await self._link_simple_objects_to_active_stages(validator)
        await self._apply_object_stages(validator)
        await self._apply_equipment(validator)
        await self._apply_work_types(validator)
        await self._apply_work_methods(validator)
        await self._apply_work_type_methods(validator)
        await self._apply_assignments(validator)
        await self._apply_object_mappings(validator)

        import_record.status = "applied"
        import_record.applied_at = sa.func.now()
        await self._session.commit()
        return import_record

    async def _load_existing(self) -> None:
        self._users = {
            u.max_user_id: u
            for u in (await self._session.execute(select(User))).scalars()
        }
        self._contractors = {
            c.code: c for c in (await self._session.execute(select(Contractor))).scalars()
        }
        self._units = {
            u.code: u for u in (await self._session.execute(select(Unit))).scalars()
        }
        self._stages = {
            s.code: s for s in (await self._session.execute(select(Stage))).scalars()
        }
        self._objects = {
            o.code: o for o in (await self._session.execute(select(Object))).scalars()
        }
        self._methods = {
            m.code: m for m in (await self._session.execute(select(WorkMethod))).scalars()
        }
        self._work_types = {
            wt.code: wt
            for wt in (await self._session.execute(select(WorkType))).scalars()
        }
        self._contracts = {
            c.code: c for c in (await self._session.execute(select(Contract))).scalars()
        }
        self._object_contracts = {
            (oc.object_id, oc.contract_id): oc
            for oc in (await self._session.execute(select(ObjectContract))).scalars()
        }

    async def _apply_users(self, validator: CatalogImportValidator) -> None:
        for row in validator.iter_rows("Users"):
            code = _normalize_text(row.get("code"))
            name = _normalize_text(row.get("name"))
            role = _normalize_text(row.get("role")) or "responsible"
            active = _normalize_bool(row.get("active"))
            if not code or not name:
                continue
            user = self._users.get(code)
            if user is None:
                user = User(max_user_id=code, full_name=name, role=role, active=active)
                self._session.add(user)
                await self._session.flush()
                self._users[code] = user
            else:
                user.full_name = name
                user.role = role
                user.active = active

    async def _apply_contractors(self, validator: CatalogImportValidator) -> None:
        for row in validator.iter_rows("Contractors"):
            code = _normalize_text(row.get("code"))
            name = _normalize_text(row.get("name"))
            active = _normalize_bool(row.get("active"))
            if not code or not name:
                continue
            contractor = await self._repo.get_or_create_contractor(code=code, name=name)
            contractor.active = active
            self._contractors[code] = contractor

    async def _apply_units(self, validator: CatalogImportValidator) -> None:
        for row in validator.iter_rows("Units"):
            code = _normalize_text(row.get("code"))
            name = _normalize_text(row.get("name"))
            symbol = _normalize_text(row.get("symbol"))
            active = _normalize_bool(row.get("active"))
            if not code or not name:
                continue
            unit = await self._repo.get_or_create_unit(code=code, name=name, symbol=symbol)
            unit.active = active
            self._units[code] = unit

    async def _apply_stages(self, validator: CatalogImportValidator) -> None:
        for row in validator.iter_rows("Stages"):
            code = _normalize_text(row.get("code"))
            name = _normalize_text(row.get("name"))
            active = _normalize_bool(row.get("active"))
            if not code or not name:
                continue
            stage = await self._repo.get_or_create_stage(code=code, name=name)
            stage.active = active
            self._stages[code] = stage

    async def _apply_objects(self, validator: CatalogImportValidator) -> None:
        for row in validator.iter_rows("Objects"):
            code = _normalize_text(row.get("code"))
            name = _normalize_text(row.get("name"))
            short_title = _normalize_text(row.get("short_title")) or None
            full_title = _normalize_text(row.get("full_title")) or None
            execution_method = _normalize_text(row.get("execution_method")) or None
            default_contractor_code = _normalize_text(row.get("default_contractor_code")) or None
            active = _normalize_bool(row.get("active"))
            if not code or not name:
                continue
            default_contractor_id = None
            if default_contractor_code:
                contractor = self._contractors.get(default_contractor_code)
                if contractor is None:
                    raise ValueError(f"Подрядчик '{default_contractor_code}' не найден")
                default_contractor_id = contractor.id
            obj = await self._repo.get_or_create_object(
                code=code,
                name=name,
                execution_method=execution_method,
                default_contractor_id=default_contractor_id,
            )
            obj.short_title = short_title
            obj.full_title = full_title
            obj.active = active
            self._objects[code] = obj

    async def _apply_object_stages(self, validator: CatalogImportValidator) -> None:
        for row in validator.iter_rows("ObjectStages"):
            object_code = _normalize_text(row.get("object_code"))
            stage_code = _normalize_text(row.get("stage_code"))
            _normalize_bool(row.get("active"))
            if not object_code or not stage_code:
                continue
            obj = self._objects.get(object_code)
            stage = self._stages.get(stage_code)
            if obj is None or stage is None:
                raise ValueError(
                    f"Связь объекта '{object_code}' и этапа '{stage_code}' невозможна"
                )
            await self._repo.ensure_object_stage(obj.id, stage.id)

    async def _link_simple_objects_to_active_stages(
        self,
        validator: CatalogImportValidator,
    ) -> None:
        """Make a two-column object import immediately usable in the form.

        A simple workbook contains no object-stage matrix.  Its active objects
        therefore inherit every active stage already maintained in the backend.
        """
        active_stages = [stage for stage in self._stages.values() if stage.active]
        if not active_stages:
            raise ValueError(
                "Нельзя импортировать простой список объектов: "
                "в базе нет активных этапов"
            )

        for row in validator.iter_rows("Objects"):
            object_code = _normalize_text(row.get("code"))
            if not object_code or not _normalize_bool(row.get("active")):
                continue
            obj = self._objects.get(object_code)
            if obj is None:
                raise ValueError(f"Объект '{object_code}' не найден после импорта")
            for stage in active_stages:
                await self._repo.ensure_object_stage(obj.id, stage.id)

    async def _apply_equipment(self, validator: CatalogImportValidator) -> None:
        for row in validator.iter_rows("Equipment"):
            code = _normalize_text(row.get("code"))
            name = _normalize_text(row.get("name"))
            default_unit_code = _normalize_text(row.get("default_unit_code"))
            active = _normalize_bool(row.get("active"))
            if not code or not name:
                continue
            unit_id = self._units.get(default_unit_code).id if default_unit_code else None
            equipment = await self._repo.get_or_create_equipment_type(
                code=code, name=name, default_unit_id=unit_id
            )
            equipment.active = active

    async def _apply_work_types(self, validator: CatalogImportValidator) -> None:
        for row in validator.iter_rows("WorkTypes"):
            code = _normalize_text(row.get("code"))
            name = _normalize_text(row.get("name"))
            default_unit_code = _normalize_text(row.get("default_unit_code"))
            active = _normalize_bool(row.get("active"))
            if not code or not name:
                continue
            unit_id = self._units.get(default_unit_code).id if default_unit_code else None
            work_type = await self._repo.get_or_create_work_type(
                code=code, name=name, default_unit_id=unit_id
            )
            work_type.active = active
            self._work_types[code] = work_type

    async def _apply_work_methods(self, validator: CatalogImportValidator) -> None:
        for row in validator.iter_rows("WorkMethods"):
            code = _normalize_text(row.get("code"))
            name = _normalize_text(row.get("name"))
            active = _normalize_bool(row.get("active"))
            if not code or not name:
                continue
            method = await self._repo.get_or_create_work_method(code=code, name=name)
            method.active = active
            self._methods[code] = method

    async def _apply_work_type_methods(self, validator: CatalogImportValidator) -> None:
        for row in validator.iter_rows("WorkTypeMethods"):
            work_type_code = _normalize_text(row.get("work_type_code"))
            work_method_code = _normalize_text(row.get("work_method_code"))
            if not work_type_code or not work_method_code:
                continue
            work_type = self._work_types.get(work_type_code)
            method = self._methods.get(work_method_code)
            if work_type is None or method is None:
                raise ValueError(
                    f"Связь вида работы '{work_type_code}' и способа '{work_method_code}' невозможна"
                )
            await self._repo.ensure_work_type_method(work_type.id, method.id)

    async def _apply_assignments(self, validator: CatalogImportValidator) -> None:
        from datetime import date as date_type

        for row in validator.iter_rows("Assignments"):
            user_code = _normalize_text(row.get("user_code"))
            object_code = _normalize_text(row.get("object_code"))
            active_from = row.get("active_from")
            active_to = row.get("active_to")
            schedule_type = _normalize_text(row.get("schedule_type")) or "daily"
            active = _normalize_bool(row.get("active"))
            if not user_code or not object_code:
                continue
            user = self._users.get(user_code)
            obj = self._objects.get(object_code)
            if user is None or obj is None:
                raise ValueError(
                    f"Назначение '{user_code}' -> '{object_code}' невозможно: пользователь или объект не найден"
                )
            if isinstance(active_from, str):
                active_from = date_type.fromisoformat(active_from)
            if isinstance(active_to, str):
                active_to = date_type.fromisoformat(active_to)
            result = await self._session.execute(
                select(ResponsibleObjectAssignment).where(
                    ResponsibleObjectAssignment.user_id == user.id,
                    ResponsibleObjectAssignment.object_id == obj.id,
                )
            )
            assignment = result.scalar_one_or_none()
            if assignment is None:
                assignment = ResponsibleObjectAssignment(
                    user_id=user.id,
                    object_id=obj.id,
                    active_from=active_from,
                    active_to=active_to,
                    schedule_type=schedule_type,
                    active=active,
                )
                self._session.add(assignment)
            else:
                assignment.active_from = active_from
                assignment.active_to = active_to
                assignment.schedule_type = schedule_type
                assignment.active = active
            await self._session.flush()

    async def _apply_object_mappings(self, validator: CatalogImportValidator) -> None:
        for row in validator.iter_rows("ObjectMappings"):
            object_code = _normalize_text(row.get("object_code"))
            if not object_code:
                continue
            obj = self._objects.get(object_code)
            if obj is None:
                raise ValueError(f"Объект '{object_code}' не найден для ObjectMappings")

            short_name = _normalize_text(row.get("short_name"))
            if short_name:
                obj.name = short_name
                if not obj.short_title:
                    obj.short_title = short_name

            contract_code = _normalize_text(row.get("contract_code"))
            if contract_code in _EMPTY_CONTRACT_MARKERS:
                continue

            full_name = _normalize_text(row.get("full_name"))
            is_primary = _normalize_bool(row.get("primary"))
            active = _normalize_bool(row.get("active"))

            # Primary contract → fill full_title on the object itself.
            if is_primary and full_name and not obj.full_title:
                obj.full_title = full_name

            contract = self._contracts.get(contract_code)
            if contract is None:
                contract = Contract(code=contract_code, full_name=full_name or contract_code)
                self._session.add(contract)
                await self._session.flush()
                self._contracts[contract_code] = contract
            else:
                if full_name:
                    contract.full_name = full_name
                contract.active = True

            existing = self._object_contracts.get((obj.id, contract.id))
            if existing is None:
                link = ObjectContract(
                    object_id=obj.id,
                    contract_id=contract.id,
                    is_primary=is_primary,
                    active=active,
                )
                self._session.add(link)
                await self._session.flush()
                self._object_contracts[(obj.id, contract.id)] = link
            else:
                existing.is_primary = is_primary
                existing.active = active


async def export_catalogs(session: AsyncSession) -> BytesIO:
    """Export current catalogs to an .xlsx workbook."""
    wb = Workbook()
    wb.remove(wb.active)

    users = (await session.execute(select(User))).scalars().all()
    contractors = (
        await session.execute(select(Contractor).where(Contractor.active))
    ).scalars().all()
    units = (await session.execute(select(Unit).where(Unit.active))).scalars().all()
    stages = (await session.execute(select(Stage).where(Stage.active))).scalars().all()
    objects = (await session.execute(select(Object))).scalars().all()
    equipment = (
        await session.execute(select(EquipmentType).where(EquipmentType.active))
    ).scalars().all()
    work_types = (
        await session.execute(select(WorkType).where(WorkType.active))
    ).scalars().all()
    methods = (
        await session.execute(select(WorkMethod).where(WorkMethod.active))
    ).scalars().all()
    object_stages = (
        await session.execute(
            select(ObjectStage).where(ObjectStage.active)
        )
    ).scalars().all()
    work_type_methods = (
        await session.execute(
            select(WorkTypeMethod).where(WorkTypeMethod.active)
        )
    ).scalars().all()
    assignments = (
        await session.execute(
            select(ResponsibleObjectAssignment).where(ResponsibleObjectAssignment.active)
        )
    ).scalars().all()
    contracts = (await session.execute(select(Contract))).scalars().all()
    object_contracts = (
        await session.execute(select(ObjectContract))
    ).scalars().all()

    code_to_contractor = {c.id: c for c in contractors}
    code_to_unit = {u.id: u for u in units}
    code_to_stage = {s.id: s for s in stages}
    code_to_object = {o.id: o for o in objects}
    code_to_work_type = {wt.id: wt for wt in work_types}
    code_to_method = {m.id: m for m in methods}
    code_to_user = {u.id: u for u in users}
    code_to_contract = {c.id: c for c in contracts}

    sheets_data = {
        "Users": [
            [u.max_user_id, u.full_name, u.role, int(u.active), ""]
            for u in users
        ],
        "Contractors": [
            [c.code, c.name, int(c.active), c.sort_order or ""]
            for c in contractors
        ],
        "Units": [
            [u.code, u.name, u.symbol, int(u.active), u.sort_order or ""]
            for u in units
        ],
        "Stages": [
            [s.code, s.name, int(s.active), s.sort_order or ""]
            for s in stages
        ],
        "Objects": [
            [
                o.code,
                o.name,
                o.short_title or "",
                o.full_title or "",
                o.execution_method or "",
                code_to_contractor.get(o.default_contractor_id, Contractor(code="")).code,
                int(o.active),
                o.sort_order or "",
            ]
            for o in objects
        ],
        "ObjectStages": [
            [
                code_to_object.get(os.object_id, Object(code="")).code,
                code_to_stage.get(os.stage_id, Stage(code="")).code,
                int(os.active),
            ]
            for os in object_stages
        ],
        "Equipment": [
            [
                e.code,
                e.name,
                code_to_unit.get(e.default_unit_id, Unit(code="")).code,
                int(e.active),
                e.sort_order or "",
            ]
            for e in equipment
        ],
        "WorkTypes": [
            [
                wt.code,
                wt.name,
                code_to_unit.get(wt.default_unit_id, Unit(code="")).code,
                int(wt.active),
                wt.sort_order or "",
            ]
            for wt in work_types
        ],
        "WorkMethods": [
            [m.code, m.name, int(m.active), m.sort_order or ""]
            for m in methods
        ],
        "WorkTypeMethods": [
            [
                code_to_work_type.get(wtm.work_type_id, WorkType(code="")).code,
                code_to_method.get(wtm.work_method_id, WorkMethod(code="")).code,
                int(wtm.active),
            ]
            for wtm in work_type_methods
        ],
        "Assignments": [
            [
                code_to_user.get(a.user_id, User(max_user_id="")).max_user_id,
                code_to_object.get(a.object_id, Object(code="")).code,
                a.active_from.isoformat() if a.active_from else "",
                a.active_to.isoformat() if a.active_to else "",
                a.schedule_type,
                int(a.active),
            ]
            for a in assignments
        ],
        "ObjectMappings": [
            [
                code_to_object.get(oc.object_id, Object(code="")).code,
                code_to_object.get(oc.object_id, Object(code="")).name,
                code_to_contract.get(oc.contract_id, Contract(code="")).code,
                code_to_contract.get(oc.contract_id, Contract(code="")).full_name,
                int(oc.is_primary),
                int(oc.active),
            ]
            for oc in object_contracts
        ],
    }

    for sheet_name in SHEET_ORDER:
        rows = sheets_data.get(sheet_name, [])
        ws = wb.create_sheet(title=sheet_name)
        ws.append(SHEET_COLUMNS[sheet_name])
        for row in rows:
            ws.append(row)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
