from __future__ import annotations

import re
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.timesheet_service import TimesheetService

_FILL_HEADER = PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid")
_FILL_MISSING = PatternFill(start_color="FFF3E0", end_color="FFF3E0", fill_type="solid")
_FILL_TOTAL = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")


def _safe_sheet_name(code: str) -> str:
    """Excel sheet names must be <=31 chars and avoid []:*?/\\."""
    cleaned = re.sub(r'[\[\]:*?/\\\\]', '-', code)
    return cleaned[:31]


def _write_header(ws, columns: list[str], row: int = 1) -> None:
    for col_idx, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=row, column=col_idx, value=col_name)
        cell.font = Font(bold=True)
        cell.fill = _FILL_HEADER
        cell.alignment = Alignment(horizontal="center")


def _autosize_columns(ws) -> None:
    for column in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)
        for cell in column:
            try:
                length = len(str(cell.value))
                if length > max_length:
                    max_length = length
            except Exception:  # noqa: BLE001
                pass
        ws.column_dimensions[column_letter].width = min(max(max_length + 2, 8), 40)


def _freeze_header_and_first_columns(ws, freeze_col: int = 3) -> None:
    ws.freeze_panes = f"{get_column_letter(freeze_col + 1)}2"


class TimesheetExcelBuilder:
    """Build an Excel workbook from timesheet data."""

    def __init__(self, service: TimesheetService) -> None:
        self._service = service

    async def build_for_object(
        self,
        object_id: int,
        date_from: str,
        date_to: str,
    ) -> BytesIO:
        data = await self._service.build(object_id, date_from, date_to)
        wb = Workbook()
        # remove default sheet
        wb.remove(wb.active)

        self._add_summary_sheet(wb, data)
        self._add_status_sheet(wb, data)
        self._add_object_sheet(wb, data)
        self._add_raw_reports_sheet(wb, data)

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    def _add_summary_sheet(self, wb: Workbook, data: dict[str, Any]) -> None:
        ws = wb.create_sheet("Общая сводка")
        ws.append(["Показатель", "Значение"])
        ws.append(["Объект", f"{data['object_code']} — {data['object_name']}"])
        ws.append(["Период", f"{data['date_from']} / {data['date_to']}"])
        ws.append(["Дней", len(data["days"])])
        ws.append(["Строк", len(data["rows"])])
        ws.append(["Пропущено дней", len(data["missing_days"])])
        _write_header(ws, ["Показатель", "Значение"])
        _autosize_columns(ws)

    def _add_status_sheet(self, wb: Workbook, data: dict[str, Any]) -> None:
        ws = wb.create_sheet("Статус отправки")
        columns = ["Дата", "Статус"]
        _write_header(ws, columns)
        for day in data["days"]:
            status = "пропущен" if day in data["missing_days"] else "сдан"
            ws.append([day, status])
        _autosize_columns(ws)
        _freeze_header_and_first_columns(ws, 0)

    def _add_object_sheet(self, wb: Workbook, data: dict[str, Any]) -> None:
        sheet_name = _safe_sheet_name(data["object_code"])
        ws = wb.create_sheet(sheet_name)
        header = ["Категория", "Показатель", "Ед.", *data["days"], "Итог", "Средн", "Макс"]
        _write_header(ws, header)
        for row in data["rows"]:
            ws.append([
                row["category"],
                row["item_name"],
                row["unit"],
                *row["values"],
                row["total"],
                row["average"],
                row["max"],
            ])
        # mark missing day columns
        for col_idx, day in enumerate(data["days"], start=4):
            if day in data["missing_days"]:
                for row_idx in range(2, ws.max_row + 1):
                    ws.cell(row=row_idx, column=col_idx).fill = _FILL_MISSING
        # style totals
        total_col = len(header) - 2
        for row_idx in range(2, ws.max_row + 1):
            ws.cell(row=row_idx, column=total_col).fill = _FILL_TOTAL
            ws.cell(row=row_idx, column=total_col).font = Font(bold=True)
        _autosize_columns(ws)
        _freeze_header_and_first_columns(ws, 3)

    def _add_raw_reports_sheet(self, wb: Workbook, data: dict[str, Any]) -> None:
        ws = wb.create_sheet("Исходные отчёты")
        ws.append(["Лист 'Исходные отчёты' заполняется в I21+ при необходимости."])
        _autosize_columns(ws)
